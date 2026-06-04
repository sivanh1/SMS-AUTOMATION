from rest_framework.decorators import api_view
from rest_framework.response import Response
from rest_framework.decorators import permission_classes
from .models import Customer
from services.google_sheets_service import get_sheet_data
from django.db.models import Q
from openpyxl import load_workbook
from rest_framework.permissions import IsAuthenticated

from rest_framework.parsers import MultiPartParser

from .serializer import XLSXUploadSerializer

from .serializer import CustomerSerializer
from services.template_engine import generate_sms_preview


#GET CUSTOMERS
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def list_customers(request):

    customers = Customer.objects.all().order_by(
        'id'
    )

    serializer = CustomerSerializer(
        customers,
        many=True
    )

    return Response(serializer.data)


#SYNC CUSTOMERS
# SYNC CUSTOMERS
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def sync_customers(request):

    sheet_id = request.data.get(
        "sheet_id"
    )

    if not sheet_id:

        return Response({

            "error":
            "Sheet ID is required"

        }, status=400)

    try:

        records = get_sheet_data(
            sheet_id
        )

        # CLEAR OLD CUSTOMERS
        Customer.objects.all().delete()

        synced_count = 0

        for row in records:

            p_id = row.get("p_id")

            if not p_id:
                continue

            # FORMAT P_ID
            if isinstance(
                p_id,
                float
            ):
                p_id = int(p_id)

            p_id = str(p_id)

            # FORMAT MOBILE NUMBER
            mobile_number = row.get(
                "mobile_number"
            )

            if mobile_number is not None:

                if isinstance(
                    mobile_number,
                    float
                ):
                    mobile_number = int(
                        mobile_number
                    )

                mobile_number = str(
                    mobile_number
                )

            # BUILD EXTRA FIELDS
            extra_fields = {}

            for key, value in row.items():

                if key not in [

                    "p_id",

                    "cust_name",

                    "mobile_number"

                ]:

                    extra_fields[key] = value

            Customer.objects.create(

                p_id=p_id,

                cust_name=row.get(
                    "cust_name"
                ),

                mobile_number=
                mobile_number,

                extra_fields=
                extra_fields,

                sheet_id=sheet_id
            )

            synced_count += 1

        return Response({

            "message":
            "Customers synchronized successfully",

            "total_synced":
            synced_count

        })

    except Exception as error:

        print(
            "SYNC ERROR:",
            error
        )

        return Response({

            "error":
            "Check Sheet ID"

        }, status=404)
    
    
# IMPORT CUSTOMERS FROM XLSX
# IMPORT CUSTOMERS FROM XLSX
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def import_customers_xlsx(request):

    serializer = XLSXUploadSerializer(
        data=request.data
    )

    if not serializer.is_valid():

        return Response(
            serializer.errors,
            status=400
        )

    file = serializer.validated_data["file"]

    try:

        workbook = load_workbook(file)

        sheet = workbook.active

        customers = []

        # GET HEADERS
        headers = []

        for cell in sheet[1]:

            headers.append(

                str(cell.value).strip()
            )

        # DATA ROWS
        for row in sheet.iter_rows(

            min_row=2,

            values_only=True
        ):

            row_data = dict(
                zip(headers, row)
            )

            p_id = row_data.get(
                "p_id"
            )

            if not p_id:
                continue

            # FORMAT P_ID
            if isinstance(
                p_id,
                float
            ):

                p_id = int(p_id)

            p_id = str(p_id)

            # CUSTOMER NAME
            cust_name = row_data.get(
                "cust_name"
            )

            # MOBILE NUMBER
            mobile_number = row_data.get(
                "mobile_number"
            )

            if mobile_number is not None:

                if isinstance(
                    mobile_number,
                    float
                ):

                    mobile_number = int(
                        mobile_number
                    )

                mobile_number = str(
                    mobile_number
                )

            # BUILD EXTRA FIELDS
            extra_fields = {}

            for key, value in row_data.items():

                if key not in [

                    "p_id",

                    "cust_name",

                    "mobile_number"

                ]:

                    # FORMAT DATES
                    if hasattr(
                        value,
                        "strftime"
                    ):

                        value = value.strftime(
                            "%Y-%m-%d"
                        )

                    extra_fields[
                        key
                    ] = value

            customers.append(

                Customer(

                    p_id=p_id,

                    cust_name=
                    cust_name,

                    mobile_number=
                    mobile_number,

                    extra_fields=
                    extra_fields
                )
            )

        # CLEAR OLD CUSTOMERS
        Customer.objects.all().delete()

        # BULK INSERT
        Customer.objects.bulk_create(
            customers
        )

        return Response({

            "message":
            "Customers imported successfully",

            "total_imported":
            len(customers)

        })

    except Exception as error:

        print(

            "XLSX IMPORT ERROR:",

            error
        )

        return Response({

            "error":
            "Failed to import XLSX file"

        }, status=500)
    
#SEARCH CUSTOMERS
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def search_customers(request):

    search = request.GET.get('search')

    customers = Customer.objects.filter(

        Q(cust_name=search) |
        Q(p_id=search)

    )

    serializer = CustomerSerializer(customers, many=True)

    return Response(serializer.data)


#GET CUSTOMER BY P_ID
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def get_customer(request):
    customer_id = request.data.get("p_id")

    if customer_id is None:
        return Response(
            {
                "error": "Please provide p_id"
            },
            status=400
        )

    try:
        customer = Customer.objects.get(
            p_id=customer_id
        )

        customer_data = CustomerSerializer(customer)

        return Response(customer_data.data)

    except Customer.DoesNotExist:
        return Response(
            {
                "error": "Customer not found"
            },
            status=404
        )
    
#PREVIEW SMS
# PREVIEW SMS
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def preview_sms(request):

    customer_id = request.data.get(
        "p_id"
    )

    template = request.data.get(
        "template"
    )

    if not customer_id or not template:

        return Response({

            "error":
            "p_id and template are required"

        }, status=400)

    try:

        customer = Customer.objects.get(
            p_id=customer_id
        )

        customer_data = {

            "p_id":
            customer.p_id,

            "cust_name":
            customer.cust_name or "Customer",

            "mobile_number":
            customer.mobile_number or "",

            **customer.extra_fields
        }

        preview_message = (
            generate_sms_preview(

                template,

                customer_data
            )
        )

        return Response({

            "preview":
            preview_message

        })

    except Customer.DoesNotExist:

        return Response({

            "error":
            "Customer not found"

        }, status=404)

    except Exception as error:

        print(
            "PREVIEW ERROR:",
            error
        )

        return Response({

            "error":
            str(error)

        }, status=500)